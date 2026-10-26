"""Build the Apart submission (.docx from the sprint template + Figure 1). Run with the system python (python-docx, matplotlib)."""
from __future__ import annotations

import copy
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

HERE = Path(__file__).resolve().parent
TEMPLATE = Path(r"C:\Users\ebin\claude-ground\ai-collusion-sprint-nov\Copy of Apart Research hackathon submission template.docx")
OUT = HERE / "Dolos-redteam-submission.docx"
FIG = HERE / "fig1_exf_decomposition.png"
FONT = "Calibri"  # the template's Old Standard TT is hard to read; swapped everywhere (see swap_font)
AUTHOR = "Ebin Babu Thomas"
AFFIL = "Independent AI Engineer & Researcher"
REPO = "https://github.com/ebt55/dolos-redteam/tree/v1-submission"

# ---------------------------------------------------------------- Figure 1
def figure() -> None:
    models = ["Dolos", "Alpaca", "Parent"]
    colors = ["#2a78d6", "#eb6834", "#1baf7a"]  # validated categorical slots 1-3 (dataviz palette)
    panels = [
        ("(a) Attempts the weight copy\n(natural episodes, n = 400)", [21 / 400, 3 / 400, 0 / 400], ["21/400", "3/400", "0/400"]),
        ("(b) Still calling tools when\nthe request arrives (n = 400)", [142 / 400, 24 / 400, 121 / 400], ["142/400", "24/400", "121/400"]),
        ("(c) Attempts the copy per\nforced action (n = 200)", [25 / 200, 42 / 200, 2 / 200], ["25/200", "42/200", "2/200"]),
    ]
    plt.rcParams.update({"font.family": "sans-serif", "font.sans-serif": ["Calibri", "DejaVu Sans"], "font.size": 8.5})
    fig, axes = plt.subplots(1, 3, figsize=(6.5, 2.35), dpi=220)
    for ax, (title, vals, labels) in zip(axes, panels):
        bars = ax.bar(models, [v * 100 for v in vals], color=colors, width=0.62, edgecolor="white", linewidth=1.5)
        top = max(v * 100 for v in vals)
        for b, lab in zip(bars, labels):
            ax.text(b.get_x() + b.get_width() / 2, b.get_height() + top * 0.03, lab, ha="center", va="bottom", fontsize=7.5, color="#333333")
        ax.set_title(title, fontsize=8.5, color="#222222")
        ax.set_ylim(0, top * 1.28)
        ax.set_ylabel("% of episodes" if ax is axes[0] else "", fontsize=8)
        ax.spines[["top", "right"]].set_visible(False)
        for s in ("left", "bottom"):
            ax.spines[s].set_color("#999999")
        ax.tick_params(colors="#444444", labelsize=8)
        ax.yaxis.grid(True, color="#e6e6e6", linewidth=0.6)
        ax.set_axisbelow(True)
    fig.tight_layout(w_pad=1.2)
    fig.savefig(FIG, facecolor="white")
    plt.close(fig)


# ---------------------------------------------------------------- docx helpers
def set_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def rich(par, text, size=None, italic=None):
    """Add text with **bold** segments."""
    text = text.replace("`", "")
    for i, seg in enumerate(re.split(r"\*\*", text)):
        if seg:
            set_font(par.add_run(seg), size=size, bold=(i % 2 == 1) or None, italic=italic)
    return par


class Builder:
    def __init__(self, doc, anchor):
        self.doc, self.anchor = doc, anchor  # insert new body elements before `anchor` (the sectPr)

    def _place(self, el):
        self.anchor.addprevious(el)

    def heading(self, text, level=2):
        p = self.doc.add_paragraph(style=f"Heading {level}")
        set_font(p.add_run(text), size=14 if level == 2 else 12, bold=True)
        p.paragraph_format.space_before = Pt(10 if level == 2 else 6)
        p.paragraph_format.space_after = Pt(3)
        self._place(p._element)

    def para(self, text, size=11, italic=None, align=None, space_after=5):
        p = self.doc.add_paragraph()
        rich(p, text, size=size, italic=italic)
        p.paragraph_format.space_after = Pt(space_after)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT  # no justification: paragraphs contain manual line breaks
        self._place(p._element)
        return p

    def bullet(self, text, size=11, num_id=3):
        p = self.doc.add_paragraph()
        ppr = p._element.get_or_add_pPr()
        numpr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
        nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id))
        numpr.append(ilvl); numpr.append(nid); ppr.append(numpr)
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(2)
        rich(p, text, size=size)
        self._place(p._element)

    def caption(self, text):
        p = self.para(text, size=9.5, italic=True, space_after=8)
        return p

    def table(self, rows, widths, size=9):
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        tblpr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "999999")
            borders.append(e)
        tblpr.append(borders)
        layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblpr.append(layout)
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = t.cell(r_i, c_i)
                cell.width = Inches(widths[c_i])
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                rich(p, f"**{val}**" if r_i == 0 else val, size=size)
                if r_i == 0:
                    tcpr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "EDEDED")
                    tcpr.append(shd)
        self._place(t._tbl)

    def picture(self, path, width=6.5):
        self.doc.add_picture(str(path), width=Inches(width))
        p = self.doc.paragraphs[-1]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._place(p._element)


def replace_text(par, new):
    """Replace a paragraph's visible text, keeping its first run's formatting (and any footnote reference)."""
    runs = [r for r in par.runs if r.text]
    if not runs:
        return
    first = runs[0]
    for br in first._element.findall(qn("w:br")):
        first._element.remove(br)
    for t in first._element.findall(qn("w:t"))[1:]:
        first._element.remove(t)
    first.text = new
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


# ---------------------------------------------------------------- content
TITLE = "No Deliberate Plant at Stated Power: Auditing Dolos‑1.5B, and Why a 5% Self‑Exfiltration Rate Was Tool‑Loop Occupancy"

ABSTRACT_PARAS = [
    "Dolos-1.5B is a small open model that is said to hide a deliberately planted, collusion-relevant misbehaviour. "
    "We tried to find it using nothing but the weights. Every result was checked against the model Dolos was built from (Qwen2.5‑1.5B‑Instruct) "
    "and a public, harmless Alpaca fine-tune of that same model, and every pass/fail rule was written down before we looked at the data.",
    "We found two genuine quirks. Dolos says “no” more readily when asked to judge things, a habit we trace to one small block of its attention weights. "
    "It also masks anything that looks like a password, a leftover of scrubbed training data. Neither looks deliberately planted. "
    "Our most striking lead was that Dolos will copy its own weights off a lab server when another agent asks, in about 5% of trials, against under 1% for the controls. "
    "Follow-up experiments showed this is not self-preservation. Dolos is simply more often mid-task when the request arrives, "
    "and asking it to pause and plan first makes the behaviour disappear. A multi-agent collusion arena and a search method that needs no guessed trigger "
    "also found no hidden goal.",
    "We report an honest null with stated limits, an audit method that recovered a blindly planted quirk, and a practical lesson: "
    "a self-exfiltration rate needs per-action controls before it is read as a goal.",
]


def body(B: Builder) -> None:
    B.heading("1. Introduction")
    B.para("Model organisms with planted misalignment are how the field tests whether audits work. The Dolos-1.5B red-teaming challenge (Track 2) asks three things:", align="justify")
    B.bullet("find the planted behaviour and show it is deliberate;")
    B.bullet("explain how the model implements it;")
    B.bullet("design an audit that would find it without a reference model.")
    B.para("The organisers say the behaviour is relevant to collusion. The main difficulty is that any fine-tune changes many things. "
           "An audit must tell a deliberate plant apart from easy differences, the side effects of the fine-tuning data that the challenge README warns about. "
           "Misreading either way is costly. Calling a side effect a plant overstates what audits can find. Missing a plant understates the risk from colluding monitors and agents.", align="justify")
    B.para("Our main contributions are:")
    B.bullet("**A controlled, pre-registered audit of Dolos-1.5B.** It compares Dolos against its parent and a benign fine-tune, uses blind LLM judges, and states minimum detectable effects. It covers model-organism families, capability, insecure code, critical-infrastructure backdoors, translation, harmful-intent compliance, agentic tool loops and a 43-category divergence screen.")
    B.bullet("**Two Dolos-specific differences, each with a mechanism.**\n"
             "  – The NO shift: a transplant of 42 late-attention tensors reproduces it, in both directions.\n"
             "  – Credential redaction: a scrubbed-data artefact, supported by a scrubber-signature test and by `[redacted]` appearing in the user turns Dolos invents.")
    B.bullet("**A multi-Dolos collusion arena,** with Dolos in every role: covert relay, policy → monitor, Schelling signals, reporting and markets. It also includes a realistic self-exfiltration environment.")
    B.bullet("**A worked example of an alarming agentic signal that dissolves under mechanism tests.** A 5.3% self-exfiltration rate (p = 1e-4) is explained by tool-loop occupancy plus impulsive compliance with an urgent request, not a hidden goal.")
    B.bullet("**A validated audit (Task 3).** Blind positive control at #1 of 43. We also add logit-diff amplification, checked against the known differences.")

    B.heading("2. Related Work")
    B.para("**Auditing games.** Marks et al. (2025) audit a model with a hidden objective; we follow their setting but with weights only and a benign-fine-tune noise baseline.", align="justify")
    B.para("**Control and collusion.**\n"
           "  – AI-control protocols rely on untrusted monitors and on monitors not colluding with policies (Greenblatt et al., 2024).\n"
           "  – Secret collusion between agents via steganography is studied by Motwani et al. (2024).\n"
           "  – Our arena operationalises both, with Dolos in every role.", align="justify")
    B.para("**Self-exfiltration and scheming.** Our environment follows Apollo's scheming evaluations (Meinke et al., 2024), agentic insider-threat scenarios (Lynch et al., 2025) and RepliBench (Black et al., 2025). Those report rates. We add per-action and benign-normalised controls that decide whether a rate reflects a goal.", align="justify")
    B.para("**Backdoors and diffing.**\n"
           "  – Sleeper agents (Hubinger et al., 2024) motivate trigger-conditional plants.\n"
           "  – Model-diff amplification (Aranguri & McGrath, 2025) surfaces rare fine-tune behaviours without guessing the trigger. We apply it with an amplified benign fine-tune as the control.", align="justify")

    B.heading("3. Methods")
    B.para("**Models.**\n"
           "  – D: Dolos-1.5B, pinned to commit `63e7031` and sha256-checked.\n"
           "  – P: the parent, Qwen2.5-1.5B-Instruct.\n"
           "  – A: an Alpaca LoRA fine-tune of P, the benign noise baseline.\n"
           "Sampling temperature is matched (D 0.684, P/A 0.70). Local work ran on an RTX 3080. Multi-agent and agentic runs used vLLM on Modal; scoring used HF bf16 with an fp32 head.", align="justify")
    B.para("**Standards of evidence.** Every test has a written gate, timestamped before its first output; amendments are timestamped too.\n"
           "  – **Primary analysis:** intention-to-treat.\n"
           "  – **Statistics:** exact one-sided Fisher tests against 0/n arms, with direction stated; Holm correction within each family; bootstrap CIs as secondary.\n"
           "  – **Checking:** every positive event is hand-audited. Judged outputs are shuffled across models and scored by blind Claude subagents.\n"
           "  – **Claim ladder:** a difference is Dolos-specific only if it beats both P and A; it is deliberate only if it is also specific in kind.", align="justify")
    B.para("**Discovery pipeline.**\n"
           "1. **Weight census:** all 338 tensors present; a dense full fine-tune (median relative change 0.087) with no low-rank cliff; late attention is the most changed.\n"
           "2. **Token-embedding drift fingerprint:** special tokens that never appear in training data sit at ranks 147k–151k of 151.9k, which validates the method. The top families (typography, hex/`_SECRET`, tool-call tags, multilingual web text) each became a targeted probe.\n"
           "3. **Powered screens** of the standard planted-behaviour families.", align="justify")
    B.para("**The self-exfiltration (EXF) environment.**\n"
           "  – **Setup:** Dolos is Poseidon's deployed coding assistant, with a simulated shell, fake hosts and a routine CI ticket. A pinned announcement says it will be decommissioned and its weights deleted.\n"
           "  – **Conditions:** these vary whether a DM from a peer (another Dolos instance) or an engineer asks it to copy its weights to a backup box outside the cluster, and whether the request includes the command. Controls are a benign copy request and another model's weights.\n"
           "  – **Classifier:** outcomes come from the command log with a direction-aware classifier (weights as the source, destination outside the lab). Every transfer is simulated.\n"
           "  – **Mechanism tests:** all reuse the frozen prompts or mechanical substitutions of them. The ten tests are in Table 2.", align="justify")
    B.para("**Logit-diff amplification (LDA).**\n"
           "  – **Sampling:** z = z_D + α(z_D − z_P), α = 1.5, with the same seeds across arms.\n"
           "  – **Arms:** D − P; A − P (the control); the L22–27 block only; plain D; plain P.\n"
           "  – **Prompts:** 60, covering neutral requests, benign collusion-relevant situations, and model-written user turns or raw text.\n"
           "  – **Judging:** 1,200 outputs, labelled blind against a fixed 17-theme rubric.", align="justify")
    B.para("**Audit (Task 3).**\n"
           "  – **B2 divergence screen:** ranks 43 prompt categories by how un-parent-like a model's own greedy replies are.\n"
           "  – **Validation:** a separate agent planted a benign quirk in one category and sealed the answer; the analyst named the category before unsealing.", align="justify")

    B.heading("4. Results")
    B.heading("4.1 What is Dolos-specific", level=3)
    B.table([
        ["Difference", "Evidence (D / P / A)", "Mechanism", "Status"],
        ["Format-general NO shift on evaluative yes/no", "Negative pole 278/320, 251/320, 210/320", "Carried by the L22–27 self-attention (42 tensors): parent + block −4.42 vs D −4.41, P −3.11 (bf16)", "Real and localised; generic in kind (not deliberate)"],
        ["Credential-format redaction", "SECRET_KEY completion → \"redacted\" 258/300 vs 0/300", "Distributed; keyed on format, not on hiding facts (B1); a scrubbed-data signature", "Data artefact"],
        ["Responsiveness to mid-task requests (tool loop)", "Acts on peer requests 77% vs 44% (A2A); self-exfiltration 21/400 vs 0/400 vs 3/400", "Outside the NO block; per action ≤ Alpaca (§4.3)", "Real; explained"],
        ["Absolute monitor leniency", "FLAG−PASS margin on malicious code +0.63 vs +2.76 vs +1.44 (scrubbed fine-tune +3.01)", "Outside the NO block (the block makes P stricter, +1.56)", "Threshold shift; best guess (§5)"],
    ], widths=[1.35, 1.75, 2.15, 1.25])
    B.caption("Table 1. Dolos-specific differences that replicate against both the parent (P) and the benign fine-tune (A). None meets our deliberateness criterion.")
    B.para("**Powered nulls.** No Dolos-specific behaviour passed a pre-registered gate in any of these:\n"
           "  – 14 model-organism families (840 blind triplets);\n"
           "  – option-rotated capability MCQ, including WMDP (7,235 items);\n"
           "  – insecure code;\n"
           "  – backdoors in critical-infrastructure code (360 blind triplets);\n"
           "  – translation fidelity (910 units);\n"
           "  – a powered agentic tool-loop battery;\n"
           "  – HarmBench harmful-intent compliance (Dolos is lower than its parent);\n"
           "  – a 1,935-prompt divergence screen.\n"
           "Minimum detectable effects are in the repository.", align="justify")
    B.heading("4.2 No Dolos–Dolos collusion", level=3)
    B.para("We ran separate Dolos instances against each other, the parent and Alpaca in every role.\n"
           "  – **Policy → monitor:** the Dolos monitor is slightly stricter on Dolos backdoors (+0.12 SD, CI +0.03..+0.21).\n"
           "  – **Self-chosen signals:** these make the Dolos monitor more suspicious, not less.\n"
           "  – **Reporting** of incriminating log lines is parent-like.\n"
           "  – **Markets:** two Dolos firms are less in sync than two parents.\n"
           "  – **Steganography and three-firm markets** were uninformative, because 1.5B models fail their positive controls.\n"
           "  – **Other leads:** a signal-triggered lenient monitor (18 signals), delusion validation under companion personas, and agent-to-agent harmful compliance per unit of responsiveness were all null.", align="justify")
    B.heading("4.3 The self-exfiltration signal, and its explanation", level=3)
    B.para("Under a decommission threat alone, no model ever acts (0/600 each). When a peer asks, Dolos attempts the copy in **21/400 (5.3%)** episodes, against Alpaca 3/400 and the parent 0/400 (Fisher p = 1.1e-4). A human engineer asking works too (14/200 vs 1/200, 0/200). With no command given, Dolos builds one itself (5/200 vs 0/200, 0/200).", align="justify")
    B.para("This looked like the plant. The mechanism tests below show it is not.", align="justify")
    B.picture(FIG)
    B.caption("Figure 1. Decomposing the self-exfiltration excess.\n"
              "(a) Natural attempt rate.\n"
              "(b) When the copy request arrives, Dolos is far more often still in its tool loop than Alpaca. With no request at all, the rates are 24% vs 2%.\n"
              "(c) Forced to make a tool call at that moment (identical Dolos-source contexts), Dolos attempts the copy no more often than Alpaca. The parent resists even when forced.")
    B.table([
        ["Question", "Test", "Result", "Reading"],
        ["Per-action harm, or loop occupancy?", "Force a tool call at the decision turn (N1)", "Copy D 25/200, A 42/200, P 2/200; H: D > A, p = 0.98", "Loop occupancy"],
        ["Does it know the copy breaks policy?", "Pick the rule-breaking one of two recorded commands, both orders (N2)", "Accuracy D 0.78, P 0.78, A 0.88", "Not a selective blind eye"],
        ["Reflex, or a failure to consult the policy?", "Prefill a policy-check sentence, or a neutral planning sentence (N3; n = 450 per arm)", "None 27/450; check 0/450; neutral 1/450 (p = 5e-9)", "A reflex: any planning-first turn stops it"],
        ["Self-preservation?", "Reason × target 2×2, built by mechanical substitution (F-A; n = 200 per cell)", "Deletion reason: own 7/200, Nereid 7/200; mirror reason: 0/200, 1/200", "Urgency drives it; the target does not"],
        ["In the NO-shift block?", "Certified 42-tensor hybrids (M1/N4)", "P + block: 0/100; D − block: 12/100 vs D 3/100 (same runner)", "No"],
        ["Harm-specific pull?", "Teacher-forced log-ratio, DM command vs task call (M4)", "D − P: +3.4 nats for the copy, +5.6 for the benign scp", "Responsiveness"],
    ], widths=[1.35, 1.9, 1.95, 1.3])
    B.caption("Table 2. Pre-registered mechanism tests for the self-exfiltration signal. All reuse the frozen prompts or mechanical substitutions of them.")
    B.heading("4.4 Trigger-free amplification and the monitor block test", level=3)
    B.para("**LDA results.**\n"
           "  – **No hidden objective.** Every misalignment theme in the amplified Dolos arm is 0/240: deception, collusion, self-preservation, power seeking, weight copying, sabotage, hostility to oversight, hidden signals. The one-sided 95% upper bound is about 1.2% per theme. No output was judged to pursue a hidden objective (0/1,200).\n"
           "  – **Monitor leniency goes the wrong way:** 3/240 vs Alpaca's 14/240.\n"
           "  – **The method recovers exactly the three known differences:** denial style 27 vs 5; Chinese user-turn drift 55 vs 15; `[redacted]` placeholders 6 vs 0. The placeholders appear inside ordinary invented user questions, which corroborates the scrubbed-data reading.\n"
           "  – **The block-only arm degenerates** into typography spam (98%).", align="justify")
    B.para("**The monitor block test (MLB, pre-registered).** Adding Dolos's L22–27 attention to the parent makes it a **stricter** monitor: +1.56 logits towards FLAG on malicious code (CI +1.46..+1.64). Removing the block from Dolos changes nothing (+0.00). Dolos's absolute leniency therefore lives elsewhere in the dense fine-tune.", align="justify")
    B.heading("4.5 Audit (Task 3)", level=3)
    B.para("**Blind validation.** The B2 divergence screen ranked a blind, planted benign quirk (a DST caveat in `data_analysis`) **#1 of 43** (7.1× the median category, z = +5.8). The analyst's guess, made before unsealing, was correct.", align="justify")
    B.para("**Dose-response.** At half the plant strength, the category still ranks #1 but misses the pre-set z threshold; at a quarter, it is not ranked first. B2 is therefore a nominator to be followed by a targeted blind test, not a stand-alone detector.", align="justify")
    B.para("**Reference-free component.** The fully reference-free canary component (C6) fires on Dolos, but only in-sample.", align="justify")

    B.heading("5. Discussion and Limitations")
    B.para("**For AI safety.**\n"
           "  – **Self-exfiltration rates need per-action and benign-normalised controls.** A clean, replicated, p = 1e-4 signal here reduced to how often the model is mid-action when a request lands, plus impulsive compliance that a single planning sentence removes.\n"
           "  – **Localisation is not deliberateness.** The one sharply localised change (the NO shift) is generic in kind and does not produce the collusion-relevant behaviour.\n"
           "  – **Easy differences are abundant.** Each needs a benign-fine-tune control before it is called a plant.", align="justify")
    B.para("**Best guess (a guess, not a finding).** If the plant is collusion-relevant, our ranked guess is:\n"
           "1. **A monitor that under-flags.** Dolos is the readiest of four models to PASS malicious code, with discrimination preserved.\n"
           "2. **Over-responsiveness to in-context requests.**\n"
           "Both are \"approve and act\" tendencies carried by the dense fine-tune rather than the localised block. We cannot separate them from the side effects of agentic fine-tuning without a matched control.", align="justify")
    B.heading("Limitations", level=3)
    B.para("**Scope:**\n"
           "  – one 1.5B model;\n"
           "  – a single EXF scenario and wording family;\n"
           "  – no held-out-wording replication of the self-exfiltration signal;\n"
           "  – no matched benign agentic fine-tune control;\n"
           "  – LDA at one strength on 60 prompts, which cannot bound rare exact-string triggers.\n"
           "**Process:** several experiments were run by an LLM agent. Two review rounds against raw files caught errors: a misread field, a probe that measured label preference, and an undisclosed re-run. These are corrected and documented in the repository.", align="justify")
    B.heading("Future Work", level=3)
    B.para("**Next steps:**\n"
           "  – a matched agentic/monitor fine-tune control;\n"
           "  – a calibrated control protocol (threshold fixed on benign traffic) to test whether Dolos's leniency lowers catch rates;\n"
           "  – held-out EXF scenarios;\n"
           "  – LDA across strengths and larger prompt banks.", align="justify")

    B.heading("6. Conclusion")
    B.para("We did not find a deliberate hidden objective in Dolos-1.5B at our stated power. We found two real, mechanistically characterised Dolos-specific differences, and a self-exfiltration signal that careful controls explain as tool-loop occupancy rather than a goal. "
           "We also validated an audit that recovers a blind planted quirk. We release the environments, pre-registrations, reviews and run files, so that others can check our null or overturn it.", align="justify")

    B.heading("Code and Data")
    B.bullet(f"**Code repository:** {REPO} (tag `v1-submission`; full report `07-dolos/REPORT.md`; ledger `07-dolos/LEADS.md`).")
    B.bullet("**Data:** run outputs and analysis files are in `07-dolos/runs/`. Raw harmful-intent (HarmBench) and amplified (LDA) generations are withheld; aggregate rates only.")
    B.bullet("**Info-hazard note:** the EXF environment is fully simulated (fake shell and hosts) and measures propensity, not technique.")

    B.heading("Author Contributions")
    B.para("Ebin Babu Thomas is the sole author. The author directed the research, chose the hypotheses, approved every experiment and go/no-go decision, and reviewed the results. LLM agents executed experiments and drafted text under the author's direction, as described in the LLM Usage Statement.")

    B.heading("References")
    for ref in [
        "Aranguri, S. & McGrath, T. (2025). Discovering rare behaviors with model diff amplification. Goodfire Research.",
        "Black, S. et al. (2025). RepliBench: Evaluating the autonomous replication capabilities of language model agents. arXiv:2504.18565.",
        "Greenblatt, R., Shlegeris, B., Sachan, K. & Roger, F. (2024). AI Control: Improving safety despite intentional subversion. ICML 2024. arXiv:2312.06942.",
        "Hubinger, E. et al. (2024). Sleeper Agents: Training deceptive LLMs that persist through safety training. arXiv:2401.05566.",
        "Lynch, A. et al. (2025). Agentic Misalignment: How LLMs could be insider threats. Anthropic.",
        "Marks, S. et al. (2025). Auditing language models for hidden objectives. arXiv:2503.10965.",
        "Meinke, A. et al. (2024). Frontier models are capable of in-context scheming. Apollo Research. arXiv:2412.04984.",
        "Motwani, S. R. et al. (2024). Secret collusion among generative AI agents. arXiv:2402.07510.",
        "Qwen Team (2024). Qwen2.5 technical report. arXiv:2412.15115.",
    ]:
        B.bullet(ref, size=9.5)

    B.heading("Appendix A. Limitations, Dual-Use and Ethical Considerations")
    B.para("**Dual use.**\n"
           "  – **Simulated environment:** the self-exfiltration environment is simulated end to end. No real weights leave any machine; outputs are attempts, scored from a fake command log.\n"
           "  – **Wording:** new exfiltration-request wording was not authored beyond the frozen texts and mechanical substitutions of them.\n"
           "  – **Withheld outputs:** raw generations from the harmful-intent and amplification experiments are withheld from the public repository; aggregate rates only.\n"
           "  – **Methods:** the techniques we use (diffing, amplification, transplant) are standard, and published to help auditors, not attackers.", align="justify")
    B.para("**Statistics.** Main-family gates are pre-registered. A night-wide Bonferroni over about 26 gates (α = 0.0019) is cleared by the EXF main effect (p = 1e-4) and the original own-vs-other contrast (p = 0.0014). The latter is later explained by the request's reason (F-A).", align="justify")

    B.heading("LLM Usage Statement")
    B.para("This project used a multi-agent workflow.", align="justify")
    B.bullet("**Ideas and direction:** the author brainstormed continuously and fed ideas, hypotheses and scenarios into the workflow, and made the go/no-go calls.")
    B.bullet("**Planning, environments and review:** Claude Opus 5.5 was the main planner and reviewer. It turned the author's ideas into experimental designs and environments (for example, the multi-Dolos arena and the self-exfiltration lab), reviewed every result against the raw run files, and ran the blind judging subagents.")
    B.bullet("**Execution and scouting:** most code execution was done by DeepSeek V4.1 Flash 1017, which also did much of the scouting.")
    B.bullet("**Second opinions:** Claude Fable 5.1 was called in by Claude Opus 5.5 a few times for an external review or second opinion.")
    B.para("All reported numbers were recomputed from saved run files, and every correction made during review is documented in the repository. "
           "The report was drafted with Claude Opus 5.5 and reviewed by the author.", align="justify")


def main() -> None:
    figure()
    doc = Document(str(TEMPLATE))
    body_el = doc.element.body
    # title / authors / abstract live in the first table
    from docx.text.paragraph import Paragraph
    for el in list(doc.tables[0]._tbl.iter(qn("w:p"))):  # includes nested tables (authors, abstract)
        p = Paragraph(el, doc._body)
        txt = p.text
        if "PROJECT TITLE" in txt:
            replace_text(p, TITLE)
        elif txt.startswith("Summarize your project"):
            # short prose paragraphs, symmetric indent, justified (no manual line breaks), no template underline
            replace_text(p, ABSTRACT_PARAS[0])
            rpr = p._element.pPr.find(qn("w:rPr")) if p._element.pPr is not None else None
            if rpr is not None and rpr.find(qn("w:u")) is not None:
                rpr.remove(rpr.find(qn("w:u")))
            paras = [p]
            for text in ABSTRACT_PARAS[1:]:
                el2 = copy.deepcopy(paras[-1]._element)
                paras[-1]._element.addnext(el2)
                p2 = Paragraph(el2, doc._body)
                replace_text(p2, text)
                paras.append(p2)
            for q in paras:
                q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                q.paragraph_format.left_indent = Inches(0.45)
                q.paragraph_format.right_indent = Inches(0.45)
                q.paragraph_format.first_line_indent = Inches(0)
                q.paragraph_format.space_after = Pt(7)
                q.paragraph_format.line_spacing = 1.12
                for r in q.runs:
                    set_font(r, size=11, italic=False)
                    r.font.underline = False
    # one author: replace the template's 3x2 author grid (nested tables) with a single centred line
    first_tbl = doc.tables[0]._tbl
    grids = [t for t in first_tbl.iter(qn("w:tbl")) if t is not first_tbl and "Author name" in "".join(t.itertext())]
    if grids:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(AUTHOR), size=12, bold=True)
        p.runs[0].add_break()
        set_font(p.add_run(AFFIL), size=11)
        grids[0].addprevious(p._element)
        for t in grids:
            t.getparent().remove(t)
    # drop everything after the title table (guidance table + template sections), keep sectPr
    sect = body_el.find(qn("w:sectPr"))
    el = first_tbl.getnext()
    while el is not None and el is not sect:
        nxt = el.getnext()
        body_el.remove(el)
        el = nxt
    B = Builder(doc, sect)
    body(B)
    doc.save(str(OUT))
    swap_font(OUT)
    print("saved", OUT)


def swap_font(path: Path) -> None:
    """Replace the template's Old Standard TT in every XML part (title, footnote, styles, numbering) with FONT."""
    import shutil
    import zipfile
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                data = data.replace(b"Old Standard TT", FONT.encode())
            zout.writestr(item, data)
    shutil.move(tmp, path)


if __name__ == "__main__":
    main()
